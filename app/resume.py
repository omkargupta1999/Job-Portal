"""Resume storage and best-effort text extraction.

Parsing is heuristic. Results are presented as *suggestions* that the user
reviews and merges explicitly; nothing is written to the profile
automatically.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import uuid
from pathlib import Path

from app.security import detect_resume_type, resolve_inside, safe_filename
from app.skills import CERTIFICATIONS, extract_skills
from app.textutil import extract_years

log = logging.getLogger(__name__)

MAX_TEXT_CHARS = 200_000


def extract_text(data: bytes, file_type: str) -> str:
    try:
        if file_type == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                return ""
            pages = []
            for page in reader.pages[:30]:
                pages.append(page.extract_text() or "")
            text = "\n".join(pages)
        elif file_type == "docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        else:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
    except Exception as exc:  # noqa: BLE001 - malformed documents are common
        log.warning("Resume text extraction failed: %s", exc)
        return ""
    text = text.replace("\x00", "")
    return text[:MAX_TEXT_CHARS]


_TITLE_RE = re.compile(
    r"\b((?:senior|sr\.?|lead|principal|staff|junior|associate)?\s*"
    r"(?:devops|devsecops|site reliability|platform|cloud(?: infrastructure| operations)?|kubernetes|openshift|"
    r"infrastructure(?: automation)?|build and release|release|ci/cd|systems|software|linux|network|security)"
    r"\s+(?:engineer|administrator|architect|consultant|specialist|lead|manager))\b",
    re.I,
)
_DEGREE_RE = re.compile(
    r"\b(b\.?\s?tech|b\.?\s?e\.?|bachelor|m\.?\s?tech|m\.?\s?sc|b\.?\s?sc|master|mba|mca|bca|ph\.?d|diploma)\b", re.I
)
_CERT_LINE_RE = re.compile(r"(certified|certification|certificate|\bCKA\b|\bCKAD\b|\bCKS\b|\bRHCSA\b|\bRHCE\b|AZ-\d{3})", re.I)
_SECTION_RE = re.compile(r"^\s*(professional summary|summary|profile|about me|objective|career summary)\s*:?\s*$", re.I)
_NEXT_SECTION_RE = re.compile(
    r"^\s*(skills|technical skills|experience|work experience|professional experience|education|certifications?|"
    r"projects|employment history|core competencies)\s*:?\s*$", re.I)


def parse_resume_text(text: str) -> dict:
    lines = [ln.strip() for ln in text.splitlines()]
    summary_lines: list[str] = []
    capture = False
    for ln in lines:
        if _SECTION_RE.match(ln):
            capture = True
            continue
        if capture and _NEXT_SECTION_RE.match(ln):
            break
        if capture and ln:
            summary_lines.append(ln)
        if len(" ".join(summary_lines)) > 1500:
            break

    titles: list[str] = []
    for m in _TITLE_RE.finditer(text):
        t = re.sub(r"\s+", " ", m.group(1)).strip().title().replace("Devops", "DevOps").replace("Devsecops", "DevSecOps")
        t = t.replace("Ci/Cd", "CI/CD").replace("Sr.", "Senior").replace("Sr ", "Senior ")
        if t not in titles:
            titles.append(t)

    skills = [s for s in extract_skills(text) if s not in CERTIFICATIONS]
    certs = []
    for ln in lines:
        if _CERT_LINE_RE.search(ln) and 3 < len(ln) < 160:
            clean = ln.lstrip("-•* ").strip()
            if clean not in certs:
                certs.append(clean)
    education = []
    for ln in lines:
        if _DEGREE_RE.search(ln) and 3 < len(ln) < 200:
            clean = ln.lstrip("-•* ").strip()
            if clean not in education:
                education.append(clean)

    years = None
    m = re.search(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs)", text, re.I)
    if m:
        value = float(m.group(1))
        if 0 < value <= 45:
            years = value
    if years is None:
        lo, _ = extract_years(text)
        years = float(lo) if lo else None

    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    return {
        "summary": " ".join(summary_lines)[:3000],
        "skills": skills[:60],
        "titles": titles[:15],
        "certifications": certs[:20],
        "education": education[:10],
        "years_experience": years,
        "email": email.group(0) if email else "",
    }


def store_resume(uploads_dir: Path, filename: str, data: bytes, max_bytes: int) -> dict:
    """Validate and save an upload. Returns metadata + extraction results."""
    file_type = detect_resume_type(filename, data, max_bytes)
    uploads_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}.{file_type}"
    path = resolve_inside(uploads_dir, stored_name)
    path.write_bytes(data)
    try:
        path.chmod(0o600)
    except OSError:
        pass
    text = extract_text(data, file_type)
    return {
        "original_filename": safe_filename(filename),
        "stored_name": stored_name,
        "file_type": file_type,
        "size_bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "extracted_text": text,
        "extracted": parse_resume_text(text) if text else {},
    }
